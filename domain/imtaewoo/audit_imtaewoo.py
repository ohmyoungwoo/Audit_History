import os
import re
import uuid
from datetime import datetime

from openpyxl import load_workbook
import pandas as pd

from fastapi import APIRouter
from fastapi import File, Depends, UploadFile, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session

from app import models, schemas
from app.dependencies import get_current_active_user, get_db
from app.utils.io import write_workbook
from app.utils.sql import safe_div
from app.core.config import cp_settings

router = APIRouter(tags=["audit"])


@router.get("/audit", response_model=list[schemas.Audit])
def get_audits(db: Session = Depends(get_db)):
    return db.query(models.Audit).all()


@router.post("/audit", response_model=schemas.Audit)
async def create_upload_files(
    db: Session = Depends(get_db),
    file: UploadFile = File(...),
    user: models.User = Depends(get_current_active_user),
):
    try:
        path = await write_workbook(file)
        wb = load_workbook(filename=path, data_only=True)

        if "Result" in wb.sheetnames:
            ws = wb["Result"]
            ws_cm = wb["Countermeasure"]
        elif "진단결과" in wb.sheetnames:
            ws = wb["진단결과"]
            ws_cm = wb["부적합-대책"]
        else:
            raise HTTPException(
                status_code=400,
                detail="Please check the sheet name. It is case sensitive and space-sensitive.",
            )

        score_df = pd.DataFrame(
            ws.iter_rows(min_row=10, max_col=6, max_row=26, values_only=True)
        )

        cm_df = pd.DataFrame(
            ws_cm.iter_rows(
                min_row=3, max_row=100, min_col=3, max_col=10, values_only=True
            )
        )

        common_df = score_df[[0, 2]].dropna()
        expertise_df = score_df[[3, 5]].dropna()
        cm_df = cm_df[cm_df[0].notna()]
        common_rate = safe_div(30, ws["B30"].value)       
        expertise_rate = safe_div(70, ws["E30"].value)        

        if type(ws["I6"].value) is datetime:
            start_date = ws["I6"].value
            end_date = None

        if type(ws["I6"].value) is str:
            date = re.findall("\d{4}.\d{2}.\d{2}", ws["I6"].value)
            if len(date) == 0:
                raise Exception(
                    "There is an error in the date format (Check Input Inside the cell I6).(e.g. 2022.12.31)"
                )
            start_date = datetime.strptime(date[0].strip(), "%Y.%m.%d")
            end_date = (
                datetime.strptime(date[1].strip(), "%Y.%m.%d")
                if len(date) > 1
                else None
            )

        type_dict = {
            "기존협력사 정기진단": "정기",
            "기존협력사 비정기진단": "비정기",
            "신규협력사진단": "신규",
        }

        if not ws["F5"].value in type_dict.keys():
            raise Exception("There is an error in the audit type format.")

        org_name_dict = cp_settings.PROD_CORP_DICT

        if not ws["I3"].value in org_name_dict.keys():
            raise Exception("There is an error in the org. name format.")

        audit_result = {
            "org_name": org_name_dict[ws["I3"].value],
            "vendor_code": ws["B4"].value,
            "vendor_name": ws["B3"].value,
            "division_name": ws["I4"].value,
            "category_code": ws["F4"].value,
            "part_name": ws["F3"].value,
            "manager_name": ws["B6"].value,
            "audit_type": type_dict[ws["F5"].value],
            "phases": re.findall("\d+", ws["I5"].value)[0],
            "common_score": ws["J10"].value,
            "expertise_score": ws["J11"].value,
            "total_score": ws["J20"].value,
            "grade": ws["H27"].value,
            "orig_file_name": file.filename,
            "file_size": os.path.getsize(path),
            "file_path": path,
            "created_by": user.username,
        }

        if not all(x is not None for x in audit_result.values()):
            raise Exception(
                "Some data was not recognized. Please check the filling form."
            )

        db_audit = models.Audit(
            start_date=start_date,
            end_date=end_date if len(date) > 1 else None,
            **audit_result,
        )

        db.add(db_audit)
        db.commit()
        db.refresh(db_audit)

        for index, row in cm_df.iterrows():
            if row[6].lower() not in ["open", "closed"]:
                db.delete(db_audit)
                db.commit()
                raise Exception("Please check status column of the audit result.")
            if row[7].lower() not in ["major", "minor", "critical"]:
                db.delete(db_audit)
                db.commit()
                raise Exception("Please check importance column of the audit result.")
            try:
                completion_date = datetime.strptime(row[3], "%Y.%m.%d")
            except:
                completion_date = None
            db_countermeasure = models.AuditCountermeasure(
                item=row[0],
                points=row[1],
                measures=row[2],
                completion_date=completion_date,
                review_result=row[5],
                status=row[6].lower(),
                importance=row[7].lower(),
                record_id=db_audit.id,
            )
            db.add(db_countermeasure)

        for index, row in common_df.iterrows():
            db_common = models.AuditCommonScore(
                order=index,
                item=row[0],
                #score=row[2] * common_rate,
                score=row[2],
                created_by=db_audit.created_by,
                record_id=db_audit.id,
            )
            db.add(db_common)

        for index, row in expertise_df.iterrows():
            db_expertise = models.AuditExpertiseScore(
                order=index,
                item=row[3],
                #score=row[5] * expertise_rate,
                score=row[5],
                created_by=db_audit.created_by,
                record_id=db_audit.id,
            )
            db.add(db_expertise)

        db.commit()
    except Exception as err:
        if os.path.isfile(path):
            os.remove(path)
        raise HTTPException(status_code=400, detail=str(err))

    return db_audit


@router.put("/audit/{item_id}", response_model=schemas.Audit)
async def update_penalty(
    item_id: int,
    file: UploadFile = File(...),
    user: models.User = Depends(get_current_active_user),
    db: Session = Depends(get_db),
):
    db_audits = read_audits_by_id(db, item_id)
    if db_audits is None:
        return None

    if db_audits.created_by != user.username:
        raise HTTPException(status_code=401)

    path = await write_workbook(file)

    try:
        wb = load_workbook(filename=path, data_only=True)

        if "Result" in wb.sheetnames:
            ws = wb["Countermeasure"]
        else:
            ws = wb["부적합-대책"]

        df = pd.DataFrame(
            ws.iter_rows(
                min_row=3, max_row=100, min_col=3, max_col=10, values_only=True
            )
        )
        df = df[df[0].notna()]

        db.query(models.AuditCountermeasure).filter(
            models.AuditCountermeasure.record_id == item_id
        ).delete()
        db.commit()

        for index, row in df.iterrows():
            if row[6].lower() not in ["open", "closed"]:
                raise Exception("Please check status column of the audit result.")
            if row[7].lower() not in ["major", "minor", "critical"]:
                raise Exception("Please check importance column of the audit result.")
            try:
                completion_date = datetime.strptime(row[3], "%Y.%m.%d")
            except:
                completion_date = None
            db_countermeasure = models.AuditCountermeasure(
                item=row[0],
                points=row[1],
                measures=row[2],
                completion_date=completion_date,
                review_result=row[5],
                status=row[6].lower(),
                importance=row[7].lower(),
                record_id=item_id,
            )
            db.add(db_countermeasure)
        db.commit()
    except Exception as err:
        if os.path.isfile(path):
            os.remove(path)
        raise HTTPException(status_code=400, detail=str(err))

    # 업데이트가 성공하면 기존 파일 삭제
    old_path = db_audits.file_path

    if os.path.isfile(old_path):
        os.remove(old_path)

    db_audits.file_path = path
    db_audits.modified_by = user.username
    db_audits.modified_date = datetime.now()
    db.commit()

    return db_audits


@router.delete("/audit/{item_id}")
def delete_audits(
    item_id: int,
    user: models.User = Depends(get_current_active_user),
    db: Session = Depends(get_db),
):
    audit = read_audits_by_id(db, item_id)

    if audit.created_by == user.username:
        db.delete(audit)
        db.commit()
    else:
        raise HTTPException(status_code=401)

    if os.path.isfile(audit.file_path):
        os.remove(audit.file_path)

    return audit


@router.get("/audit/file/{item_id}")
async def download_file(item_id: int, db: Session = Depends(get_db)):
    file = read_audits_by_id(db, item_id)
    return FileResponse(path=os.path.join(file.file_path))


def read_audits_by_id(db: Session, item_id: int):
    return db.query(models.Audit).filter(models.Audit.id == item_id).first()


async def write_workbook(file):
    try:
        basename, ext = os.path.splitext(file.filename)
        path = os.path.join("./uploads/audits", str(uuid.uuid1()) + ext)

        contents = await file.read()
        with open(path, "wb") as fp:
            fp.write(contents)
        return path
    except:
        raise HTTPException(status_code=400, detail="Upload failed.")
